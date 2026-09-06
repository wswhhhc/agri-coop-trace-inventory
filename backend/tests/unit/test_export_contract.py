from datetime import date
from io import BytesIO
from uuid import uuid4

import pytest
from app.core.auth.context import AuthContext
from app.core.exceptions import AppException
from app.schemas.export import ExportFilters, ExportTaskCreate, ReportType
from app.services.export_policy import require_export
from app.services.export_workbook import build_alert_workbook, build_inventory_workbook
from openpyxl import load_workbook


def _context(*, role_code: str, permissions: set[str]) -> AuthContext:
    return AuthContext(
        user_id=uuid4(),
        username="export-user",
        real_name="导出用户",
        role_code=role_code,
        permission_codes=frozenset(permissions),
        cooperative_id=uuid4(),
        warehouse_ids=None,
        session_id="session",
        token_id="token",
    )


def test_export_schema_accepts_camel_case_and_rejects_range_over_366_days() -> None:
    payload = ExportTaskCreate(
        reportType="INVENTORY_DETAIL",
        filters={
            "warehouseId": str(uuid4()),
            "startDate": "2026-09-01",
            "endDate": "2026-09-30",
        },
    )

    assert payload.report_type is ReportType.INVENTORY_DETAIL
    assert payload.filters.start_date == date(2026, 9, 1)

    with pytest.raises(ValueError, match="366"):
        ExportFilters(startDate="2025-01-01", endDate="2026-01-02")


def test_export_schema_defaults_filters() -> None:
    payload = ExportTaskCreate(reportType=ReportType.ALERT_DETAIL)

    assert payload.filters.start_date is None
    assert payload.filters.end_date is None


def test_export_policy_requires_cooperative_admin_and_report_permission() -> None:
    require_export(
        _context(role_code="COOPERATIVE_ADMIN", permissions={"inventory:read"}),
        ReportType.INVENTORY_DETAIL,
    )
    require_export(
        _context(role_code="COOPERATIVE_ADMIN", permissions={"alert:read"}),
        ReportType.ALERT_DETAIL,
    )

    with pytest.raises(AppException) as error:
        require_export(
            _context(role_code="WAREHOUSE_STAFF", permissions={"inventory:read"}),
            ReportType.INVENTORY_DETAIL,
        )
    assert error.value.status_code == 403


def test_inventory_workbook_has_fixed_headers_and_numeric_columns() -> None:
    output = BytesIO()
    build_inventory_workbook(
        [
            {
                "date": date(2026, 9, 1),
                "warehouse": "中心仓",
                "batch_no": "B-001",
                "product_code": "RICE",
                "product_name": "优质粳米",
                "unit": "千克",
                "quantity": 12.5,
                "locked_quantity": 1.5,
                "available_quantity": 11.0,
            }
        ],
        output,
    )

    sheet = load_workbook(output).active
    assert sheet.title == "库存明细"
    assert [cell.value for cell in sheet[1]] == [
        "日期",
        "仓库",
        "批次号",
        "产品编码",
        "产品名称",
        "单位",
        "库存数量",
        "锁定数量",
        "可用数量",
    ]
    assert sheet.cell(2, 1).value.date() == date(2026, 9, 1)
    assert sheet.cell(2, 7).value == 12.5
    assert sheet.freeze_panes == "A2"


def test_alert_workbook_contains_status_and_evidence_columns() -> None:
    output = BytesIO()
    build_alert_workbook(
        [
            {
                "detected_at": date(2026, 9, 1),
                "alert_type": "LOW_STOCK",
                "severity": "HIGH",
                "status": "PENDING",
                "warehouse": "中心仓",
                "product": "优质粳米",
                "batch_no": "B-001",
                "title": "库存不足",
                "message": "库存低于安全库存",
                "resolved_at": None,
            }
        ],
        output,
    )

    sheet = load_workbook(output).active
    assert sheet.title == "预警明细"
    assert [cell.value for cell in sheet[1]] == [
        "检测时间",
        "预警类型",
        "级别",
        "状态",
        "仓库",
        "产品",
        "批次号",
        "标题",
        "说明",
        "处理时间",
    ]
    assert sheet.cell(2, 3).value == "HIGH"
