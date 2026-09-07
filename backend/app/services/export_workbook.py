from __future__ import annotations

from collections.abc import Mapping, Sequence
from datetime import date, datetime
from typing import Any, BinaryIO

from openpyxl import Workbook  # type: ignore[import-untyped]
from openpyxl.styles import Font  # type: ignore[import-untyped]
from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]

INVENTORY_HEADERS = (
    "日期",
    "仓库",
    "批次号",
    "产品编码",
    "产品名称",
    "单位",
    "库存数量",
    "锁定数量",
    "可用数量",
)
ALERT_HEADERS = (
    "检测时间",
    "预警类型",
    "级别",
    "状态",
    "仓库",
    "产品",
    "批次号",
    "标题",
    "说明",
    "处理时间",
)


def build_inventory_workbook(
    rows: Sequence[Mapping[str, Any]], output: BinaryIO
) -> None:
    _write_workbook(
        "库存明细",
        INVENTORY_HEADERS,
        rows,
        (
            "date",
            "warehouse",
            "batch_no",
            "product_code",
            "product_name",
            "unit",
            "quantity",
            "locked_quantity",
            "available_quantity",
        ),
        output,
    )


def build_alert_workbook(rows: Sequence[Mapping[str, Any]], output: BinaryIO) -> None:
    _write_workbook(
        "预警明细",
        ALERT_HEADERS,
        rows,
        (
            "detected_at",
            "alert_type",
            "severity",
            "status",
            "warehouse",
            "product",
            "batch_no",
            "title",
            "message",
            "resolved_at",
        ),
        output,
    )


def _write_workbook(
    title: str,
    headers: Sequence[str],
    rows: Sequence[Mapping[str, Any]],
    fields: Sequence[str],
    output: BinaryIO,
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title
    sheet.append(list(headers))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append([_to_excel_value(row.get(field)) for field in fields])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column_index, values in enumerate(sheet.iter_cols(), start=1):
        length = max(len(str(value.value or "")) for value in values)
        sheet.column_dimensions[get_column_letter(column_index)].width = min(max(length + 2, 10), 30)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, (date, datetime)):
                cell.number_format = "yyyy-mm-dd" if isinstance(cell.value, date) and not isinstance(cell.value, datetime) else "yyyy-mm-dd hh:mm:ss"
    workbook.save(output)


def _to_excel_value(value: Any) -> Any:
    """移除 Excel 不支持的 datetime 时区信息，保留原始时间值。"""
    if isinstance(value, datetime) and value.tzinfo is not None:
        return value.replace(tzinfo=None)
    return value


__all__ = ["build_alert_workbook", "build_inventory_workbook"]
